from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import EmployeeReport


HEADERS = [
    "Report Date",
    "Updated At",
    "Project Code",
    "Employee Name",
    "Telegram Username",
    "User ID",
    "Chat ID",
    "Message ID",
    "Source",
    "Telegram File ID",
    "Transcript",
    "Fact Today",
    "Plan Tomorrow",
    "Parse Status",
]


class ExcelReportStore:
    def __init__(self, workbook_path: Path) -> None:
        self._workbook_path = workbook_path
        self._workbook_path.parent.mkdir(parents=True, exist_ok=True)

    def upsert_report(self, report: EmployeeReport) -> None:
        workbook = self._load_or_create_workbook()
        sheet = workbook["Reports"]
        header_map = self._ensure_headers(sheet)
        row_index = self._find_row(sheet, header_map, report.report_date, report.user_id, report.project_code)
        values = {
            "Report Date": report.report_date,
            "Updated At": report.updated_at,
            "Project Code": report.project_code,
            "Employee Name": report.employee_name,
            "Telegram Username": report.telegram_username,
            "User ID": report.user_id,
            "Chat ID": report.chat_id,
            "Message ID": report.message_id,
            "Source": report.source,
            "Telegram File ID": report.telegram_file_id,
            "Transcript": report.transcript,
            "Fact Today": report.fact_today,
            "Plan Tomorrow": report.plan_tomorrow,
            "Parse Status": report.parse_status,
        }

        if row_index is None:
            row_index = sheet.max_row + 1

        for header, value in values.items():
            sheet.cell(row=row_index, column=header_map[header], value=value)

        self._finalize_sheet(sheet)
        workbook.save(self._workbook_path)

    def get_reports_for_date(
        self,
        report_date: str,
        project_code: str | None = None,
    ) -> list[dict[str, str]]:
        workbook = self._load_or_create_workbook()
        sheet = workbook["Reports"]
        header_map = self._ensure_headers(sheet)
        reports: list[dict[str, str]] = []

        for row_index in range(2, sheet.max_row + 1):
            row_date = str(sheet.cell(row=row_index, column=header_map["Report Date"]).value or "")
            row_project = str(sheet.cell(row=row_index, column=header_map["Project Code"]).value or "")
            if row_date != report_date:
                continue
            if project_code is not None and row_project != project_code:
                continue

            reports.append(
                {
                    "report_date": row_date,
                    "updated_at": str(sheet.cell(row=row_index, column=header_map["Updated At"]).value or ""),
                    "project_code": row_project,
                    "employee_name": str(sheet.cell(row=row_index, column=header_map["Employee Name"]).value or ""),
                    "telegram_username": str(sheet.cell(row=row_index, column=header_map["Telegram Username"]).value or ""),
                    "user_id": str(sheet.cell(row=row_index, column=header_map["User ID"]).value or ""),
                    "source": str(sheet.cell(row=row_index, column=header_map["Source"]).value or ""),
                    "fact_today": str(sheet.cell(row=row_index, column=header_map["Fact Today"]).value or ""),
                    "plan_tomorrow": str(sheet.cell(row=row_index, column=header_map["Plan Tomorrow"]).value or ""),
                    "parse_status": str(sheet.cell(row=row_index, column=header_map["Parse Status"]).value or ""),
                }
            )

        return reports

    def export_reports(
        self,
        output_path: Path,
        project_code: str | None = None,
    ) -> Path:
        source_workbook = self._load_or_create_workbook()
        source_sheet = source_workbook["Reports"]
        header_map = self._ensure_headers(source_sheet)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        export_workbook = Workbook()
        export_sheet = export_workbook.active
        export_sheet.title = "Reports"
        export_sheet.append(HEADERS)

        for row_index in range(2, source_sheet.max_row + 1):
            row_project = str(source_sheet.cell(row=row_index, column=header_map["Project Code"]).value or "")
            if project_code is not None and row_project != project_code:
                continue

            export_sheet.append(
                [
                    source_sheet.cell(row=row_index, column=header_map[header]).value or ""
                    for header in HEADERS
                ]
            )

        self._finalize_sheet(export_sheet)
        export_workbook.save(output_path)
        return output_path

    def _load_or_create_workbook(self):
        if self._workbook_path.exists():
            workbook = load_workbook(self._workbook_path)
            sheet = workbook["Reports"]
            self._ensure_headers(sheet)
            self._finalize_sheet(sheet)
            return workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reports"
        sheet.append(HEADERS)
        self._finalize_sheet(sheet)
        workbook.save(self._workbook_path)
        return workbook

    def _ensure_headers(self, sheet: Worksheet) -> dict[str, int]:
        existing_header_map: dict[str, int] = {}
        for column_index in range(1, sheet.max_column + 1):
            header_value = str(sheet.cell(row=1, column=column_index).value or "").strip()
            if header_value:
                existing_header_map[header_value] = column_index

        for header in HEADERS:
            if header in existing_header_map:
                continue
            column_index = sheet.max_column + 1
            sheet.cell(row=1, column=column_index, value=header)
            existing_header_map[header] = column_index

        return existing_header_map

    @staticmethod
    def _finalize_sheet(sheet: Worksheet) -> None:
        last_column_letter = get_column_letter(sheet.max_column)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{last_column_letter}1"

    @staticmethod
    def _find_row(
        sheet: Worksheet,
        header_map: dict[str, int],
        report_date: str,
        user_id: int,
        project_code: str,
    ) -> int | None:
        for row_index in range(2, sheet.max_row + 1):
            row_date = str(sheet.cell(row=row_index, column=header_map["Report Date"]).value or "")
            row_user_id = str(sheet.cell(row=row_index, column=header_map["User ID"]).value or "")
            row_project = str(sheet.cell(row=row_index, column=header_map["Project Code"]).value or "")
            if row_date == report_date and row_user_id == str(user_id) and row_project == project_code:
                return row_index
        return None
