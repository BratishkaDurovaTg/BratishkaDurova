from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from psycopg import connect
from psycopg.rows import dict_row

from .excel_store import HEADERS
from .models import EmployeeReport
from .project_store import DEFAULT_PROJECT_CODES, ProjectStore


TIMESTAMP_FORMAT = "%Y-%m-%d %H:%M:%S"


class PostgresProjectStore:
    def __init__(self, database_url: str, snapshot_path: Path | None = None) -> None:
        self._database_url = database_url
        self._snapshot_path = snapshot_path
        _ensure_schema(self._database_url)
        self._bootstrap_projects()

    def list_projects(self) -> list[str]:
        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT code FROM projects ORDER BY code")
                return [row["code"] for row in cur.fetchall()]

    def add_project(self, project_code: str) -> tuple[str, bool]:
        normalized_code = ProjectStore.normalize_project_code(project_code)
        if not normalized_code:
            raise ValueError(
                "Код проекта должен содержать только латинские буквы, цифры, '-' или '_' и быть не длиннее 20 символов."
            )

        created = False
        with connect(self._database_url) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    INSERT INTO projects (code)
                    VALUES (%s)
                    ON CONFLICT (code) DO NOTHING
                    """,
                    (normalized_code,),
                )
                created = cur.rowcount > 0

        self._sync_snapshot()
        return normalized_code, created

    def _bootstrap_projects(self) -> None:
        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT COUNT(*) AS total FROM projects")
                total = int(cur.fetchone()["total"])
                if total > 0:
                    self._sync_snapshot()
                    return

                for code in _load_project_codes(self._snapshot_path):
                    cur.execute(
                        """
                        INSERT INTO projects (code)
                        VALUES (%s)
                        ON CONFLICT (code) DO NOTHING
                        """,
                        (code,),
                    )

        self._sync_snapshot()

    def _sync_snapshot(self) -> None:
        if self._snapshot_path is None:
            return

        self._snapshot_path.parent.mkdir(parents=True, exist_ok=True)
        self._snapshot_path.write_text(
            json.dumps(self.list_projects(), ensure_ascii=True, indent=2) + "\n",
            encoding="utf-8",
        )


class PostgresReportStore:
    def __init__(self, database_url: str, snapshot_path: Path | None = None) -> None:
        self._database_url = database_url
        self._snapshot_path = snapshot_path
        _ensure_schema(self._database_url)
        self._bootstrap_reports()

    def upsert_report(self, report: EmployeeReport) -> None:
        report_date = _parse_report_date(report.report_date)
        updated_at = _parse_updated_at(report.updated_at)

        with connect(self._database_url) as conn:
            with conn.cursor() as cur:
                _ensure_project(cur, report.project_code)
                cur.execute(
                    """
                    INSERT INTO reports (
                        report_date,
                        updated_at,
                        project_code,
                        employee_name,
                        telegram_username,
                        user_id,
                        chat_id,
                        message_id,
                        source,
                        telegram_file_id,
                        transcript,
                        fact_today,
                        plan_tomorrow,
                        parse_status
                    )
                    VALUES (
                        %(report_date)s,
                        %(updated_at)s,
                        %(project_code)s,
                        %(employee_name)s,
                        %(telegram_username)s,
                        %(user_id)s,
                        %(chat_id)s,
                        %(message_id)s,
                        %(source)s,
                        %(telegram_file_id)s,
                        %(transcript)s,
                        %(fact_today)s,
                        %(plan_tomorrow)s,
                        %(parse_status)s
                    )
                    ON CONFLICT (report_date, user_id, project_code) DO UPDATE SET
                        updated_at = EXCLUDED.updated_at,
                        employee_name = EXCLUDED.employee_name,
                        telegram_username = EXCLUDED.telegram_username,
                        chat_id = EXCLUDED.chat_id,
                        message_id = EXCLUDED.message_id,
                        source = EXCLUDED.source,
                        telegram_file_id = EXCLUDED.telegram_file_id,
                        transcript = EXCLUDED.transcript,
                        fact_today = EXCLUDED.fact_today,
                        plan_tomorrow = EXCLUDED.plan_tomorrow,
                        parse_status = EXCLUDED.parse_status
                    """,
                    {
                        "report_date": report_date,
                        "updated_at": updated_at,
                        "project_code": report.project_code,
                        "employee_name": report.employee_name,
                        "telegram_username": report.telegram_username,
                        "user_id": report.user_id,
                        "chat_id": report.chat_id,
                        "message_id": report.message_id,
                        "source": report.source,
                        "telegram_file_id": report.telegram_file_id,
                        "transcript": report.transcript,
                        "fact_today": report.fact_today,
                        "plan_tomorrow": report.plan_tomorrow,
                        "parse_status": report.parse_status,
                    },
                )

        self._refresh_snapshot()

    def get_reports_for_date(
        self,
        report_date: str,
        project_code: str | None = None,
    ) -> list[dict[str, str]]:
        query = """
            SELECT
                report_date,
                updated_at,
                project_code,
                employee_name,
                telegram_username,
                user_id,
                source,
                fact_today,
                plan_tomorrow,
                parse_status
            FROM reports
            WHERE report_date = %s
        """
        params: list[Any] = [_parse_report_date(report_date)]
        if project_code is not None:
            query += " AND project_code = %s"
            params.append(project_code)
        query += " ORDER BY updated_at DESC, employee_name ASC, user_id ASC"

        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
                rows = cur.fetchall()

        return [_row_to_report_dict(row) for row in rows]

    def list_recent_reports(
        self,
        limit: int | None = 20,
        project_code: str | None = None,
    ) -> list[dict[str, str]]:
        query = """
            SELECT
                report_date,
                updated_at,
                project_code,
                employee_name,
                telegram_username,
                user_id,
                chat_id,
                message_id,
                source,
                telegram_file_id,
                transcript,
                fact_today,
                plan_tomorrow,
                parse_status
            FROM reports
        """
        params: list[Any] = []
        if project_code is not None:
            query += " WHERE project_code = %s"
            params.append(project_code)
        query += " ORDER BY report_date DESC, updated_at DESC, employee_name ASC"
        if limit is not None:
            query += " LIMIT %s"
            params.append(limit)

        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
                return [_row_to_report_dict(row) for row in cur.fetchall()]

    def get_user_reports(
        self,
        user_id: int,
        limit: int | None = 20,
    ) -> list[dict[str, str]]:
        query = """
            SELECT
                report_date,
                updated_at,
                project_code,
                employee_name,
                telegram_username,
                user_id,
                chat_id,
                message_id,
                source,
                telegram_file_id,
                transcript,
                fact_today,
                plan_tomorrow,
                parse_status
            FROM reports
            WHERE user_id = %s
            ORDER BY report_date DESC, updated_at DESC, project_code ASC
        """
        params: list[Any] = [user_id]
        if limit is not None:
            query += " LIMIT %s"
            params.append(limit)

        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
                return [_row_to_report_dict(row) for row in cur.fetchall()]

    def export_reports(
        self,
        output_path: Path,
        project_code: str | None = None,
    ) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reports"
        sheet.append(HEADERS)

        query = """
            SELECT
                report_date,
                updated_at,
                project_code,
                employee_name,
                telegram_username,
                user_id,
                chat_id,
                message_id,
                source,
                telegram_file_id,
                transcript,
                fact_today,
                plan_tomorrow,
                parse_status
            FROM reports
        """
        params: list[Any] = []
        if project_code is not None:
            query += " WHERE project_code = %s"
            params.append(project_code)
        query += " ORDER BY report_date DESC, updated_at DESC, project_code ASC, employee_name ASC"

        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute(query, params)
                for row in cur.fetchall():
                    sheet.append(
                        [
                            row["report_date"].isoformat(),
                            row["updated_at"].strftime(TIMESTAMP_FORMAT),
                            row["project_code"],
                            row["employee_name"],
                            row["telegram_username"],
                            row["user_id"],
                            row["chat_id"],
                            row["message_id"],
                            row["source"],
                            row["telegram_file_id"],
                            row["transcript"],
                            row["fact_today"],
                            row["plan_tomorrow"],
                            row["parse_status"],
                        ]
                    )

        _finalize_sheet(sheet)
        workbook.save(output_path)
        return output_path

    def _bootstrap_reports(self) -> None:
        with connect(self._database_url, row_factory=dict_row) as conn:
            with conn.cursor() as cur:
                cur.execute("SELECT COUNT(*) AS total FROM reports")
                total = int(cur.fetchone()["total"])
                if total > 0:
                    self._refresh_snapshot()
                    return

        reports = _load_reports_from_workbook(self._snapshot_path)
        if not reports:
            self._refresh_snapshot()
            return

        with connect(self._database_url) as conn:
            with conn.cursor() as cur:
                for report in reports:
                    _ensure_project(cur, report.project_code)
                    cur.execute(
                        """
                        INSERT INTO reports (
                            report_date,
                            updated_at,
                            project_code,
                            employee_name,
                            telegram_username,
                            user_id,
                            chat_id,
                            message_id,
                            source,
                            telegram_file_id,
                            transcript,
                            fact_today,
                            plan_tomorrow,
                            parse_status
                        )
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        ON CONFLICT (report_date, user_id, project_code) DO NOTHING
                        """,
                        (
                            _parse_report_date(report.report_date),
                            _parse_updated_at(report.updated_at),
                            report.project_code,
                            report.employee_name,
                            report.telegram_username,
                            report.user_id,
                            report.chat_id,
                            report.message_id,
                            report.source,
                            report.telegram_file_id,
                            report.transcript,
                            report.fact_today,
                            report.plan_tomorrow,
                            report.parse_status,
                        ),
                    )

        self._refresh_snapshot()

    def _refresh_snapshot(self) -> None:
        if self._snapshot_path is None:
            return
        self.export_reports(self._snapshot_path)


def _ensure_schema(database_url: str) -> None:
    with connect(database_url) as conn:
        with conn.cursor() as cur:
            # Serialize bootstrap so bot and web can start together safely.
            cur.execute("SELECT pg_advisory_xact_lock(%s)", (428114907511,))
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS projects (
                    code TEXT PRIMARY KEY,
                    created_at TIMESTAMP WITHOUT TIME ZONE NOT NULL DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS reports (
                    report_date DATE NOT NULL,
                    updated_at TIMESTAMP WITHOUT TIME ZONE NOT NULL,
                    project_code TEXT NOT NULL REFERENCES projects(code) ON DELETE RESTRICT,
                    employee_name TEXT NOT NULL,
                    telegram_username TEXT NOT NULL DEFAULT '',
                    user_id BIGINT NOT NULL,
                    chat_id BIGINT NOT NULL,
                    message_id BIGINT NOT NULL,
                    source TEXT NOT NULL,
                    telegram_file_id TEXT NOT NULL DEFAULT '',
                    transcript TEXT NOT NULL,
                    fact_today TEXT NOT NULL DEFAULT '',
                    plan_tomorrow TEXT NOT NULL DEFAULT '',
                    parse_status TEXT NOT NULL,
                    PRIMARY KEY (report_date, user_id, project_code)
                )
                """
            )
            cur.execute(
                """
                CREATE INDEX IF NOT EXISTS reports_report_date_project_idx
                ON reports (report_date, project_code)
                """
            )


def _ensure_project(cur, project_code: str) -> None:
    cur.execute(
        """
        INSERT INTO projects (code)
        VALUES (%s)
        ON CONFLICT (code) DO NOTHING
        """,
        (project_code,),
    )


def _load_project_codes(projects_path: Path | None) -> list[str]:
    if projects_path is None or not projects_path.exists():
        return list(DEFAULT_PROJECT_CODES)

    try:
        payload = json.loads(projects_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return list(DEFAULT_PROJECT_CODES)

    if not isinstance(payload, list):
        return list(DEFAULT_PROJECT_CODES)

    project_codes: list[str] = []
    for item in payload:
        if not isinstance(item, str):
            continue
        normalized = ProjectStore.normalize_project_code(item)
        if normalized and normalized not in project_codes:
            project_codes.append(normalized)

    return project_codes or list(DEFAULT_PROJECT_CODES)


def _load_reports_from_workbook(workbook_path: Path | None) -> list[EmployeeReport]:
    if workbook_path is None or not workbook_path.exists():
        return []

    try:
        workbook = load_workbook(workbook_path, read_only=True)
    except OSError:
        return []

    if "Reports" not in workbook.sheetnames:
        workbook.close()
        return []

    sheet = workbook["Reports"]
    rows = sheet.iter_rows(values_only=True)
    header_row = next(rows, None)
    if header_row is None:
        workbook.close()
        return []

    header_map = {
        str(value).strip(): index
        for index, value in enumerate(header_row)
        if str(value or "").strip()
    }

    reports: list[EmployeeReport] = []
    for row in rows:
        report = _build_employee_report_from_row(row, header_map)
        if report is not None:
            reports.append(report)

    workbook.close()
    return reports


def _build_employee_report_from_row(
    row: tuple[Any, ...],
    header_map: dict[str, int],
) -> EmployeeReport | None:
    def get_value(header: str, default: Any = "") -> Any:
        index = header_map.get(header)
        if index is None or index >= len(row):
            return default
        value = row[index]
        return default if value is None else value

    try:
        project_code = ProjectStore.normalize_project_code(str(get_value("Project Code", "")))
        report_date = _normalize_report_date(get_value("Report Date", ""))
        updated_at = _normalize_updated_at(get_value("Updated At", ""))
        user_id = _to_int(get_value("User ID", 0))
        chat_id = _to_int(get_value("Chat ID", 0))
        message_id = _to_int(get_value("Message ID", 0))
    except (TypeError, ValueError):
        return None

    if not project_code or not report_date or user_id == 0:
        return None

    return EmployeeReport(
        report_date=report_date,
        updated_at=updated_at,
        project_code=project_code,
        employee_name=str(get_value("Employee Name", "")),
        telegram_username=str(get_value("Telegram Username", "")),
        user_id=user_id,
        chat_id=chat_id,
        message_id=message_id,
        source=str(get_value("Source", "")),
        telegram_file_id=str(get_value("Telegram File ID", "")),
        transcript=str(get_value("Transcript", "")),
        fact_today=str(get_value("Fact Today", "")),
        plan_tomorrow=str(get_value("Plan Tomorrow", "")),
        parse_status=str(get_value("Parse Status", "transcript_only")),
    )


def _normalize_report_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _parse_report_date(str(value)).isoformat()


def _normalize_updated_at(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime(TIMESTAMP_FORMAT)
    return _parse_updated_at(str(value)).strftime(TIMESTAMP_FORMAT)


def _parse_report_date(value: str) -> date:
    return date.fromisoformat((value or "").strip())


def _parse_updated_at(value: str) -> datetime:
    raw_value = (value or "").strip()
    if not raw_value:
        raise ValueError("Updated At value is required for PostgreSQL storage.")

    try:
        return datetime.fromisoformat(raw_value)
    except ValueError:
        return datetime.strptime(raw_value, TIMESTAMP_FORMAT)


def _to_int(value: Any) -> int:
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    return int(str(value).strip())


def _finalize_sheet(sheet) -> None:
    from openpyxl.utils import get_column_letter

    last_column_letter = get_column_letter(sheet.max_column)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last_column_letter}1"


def _row_to_report_dict(row: dict[str, Any]) -> dict[str, str]:
    return {
        "report_date": row["report_date"].isoformat(),
        "updated_at": row["updated_at"].strftime(TIMESTAMP_FORMAT),
        "project_code": row["project_code"],
        "employee_name": row["employee_name"],
        "telegram_username": row["telegram_username"],
        "user_id": str(row["user_id"]),
        "chat_id": str(row.get("chat_id", "")),
        "message_id": str(row.get("message_id", "")),
        "source": row["source"],
        "telegram_file_id": row.get("telegram_file_id", ""),
        "transcript": row.get("transcript", ""),
        "fact_today": row["fact_today"],
        "plan_tomorrow": row["plan_tomorrow"],
        "parse_status": row["parse_status"],
    }
